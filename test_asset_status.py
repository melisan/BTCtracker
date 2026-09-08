import copy
import ast
import os
from pathlib import Path
import sqlite3
import unittest
from io import BytesIO
from datetime import date
import openpyxl
from unittest.mock import patch

from cryptography.fernet import Fernet
from flask import Flask
from asset_status import create_asset_blueprint, PRIVATE_TAB, AUTH_TAB, setup_capability, validate_document


SAMPLE = {"accounts":[{"id":"sample-1", "in_total":True, "category":"예적금", "name":"가상 이름", "description":"가상 예금",
          "account":"001-000-000", "amount":100000, "interest":1234, "notes":"테스트", "maturity":"2026-12-31"}], "other":"가상 기타 내용", "other_assets":[], "total_scope":"c1"}


def fixture():
    db = sqlite3.connect(":memory:", check_same_thread=False)
    db.row_factory = sqlite3.Row
    db.execute("CREATE TABLE tab_notes (id INTEGER PRIMARY KEY, tab TEXT, content TEXT, pinned BOOLEAN DEFAULT FALSE, created_at TEXT)")
    class Connection:
        def commit(self): db.commit()
        def close(self): db.rollback()
    conn = Connection()
    def query(_, sql, args=()):
        if "pg_advisory" in sql: return None
        return db.execute(sql.replace("%s", "?"), args)
    app = Flask(__name__)
    app.register_blueprint(create_asset_blueprint(lambda:conn, query))
    return app, db


class SecurityTests(unittest.TestCase):
    def setUp(self):
        self.env = patch.dict(os.environ, {"ASSET_STATUS_KEY":Fernet.generate_key().decode(), "ADMIN_PASSWORD":"synthetic-test-only-password"})
        self.env.start()
        self.app, self.db = fixture()
        self.client = self.app.test_client()
        self.headers = {"X-Asset-Request":"1"}
        self.client.post("/asset-status/setup",json={"password":"synthetic-test-only-password"},headers={**self.headers,"X-Asset-Setup":setup_capability()})
    def tearDown(self):
        self.db.close(); self.env.stop()
    def unlock(self):
        response = self.client.post("/asset-status/unlock", json={"password":"synthetic-test-only-password"}, headers=self.headers)
        if response.status_code == 200:
            self.headers["X-Asset-Token"] = response.json["token"]
        return response
    def test_requires_password_for_read_and_write(self):
        self.assertEqual(self.client.get("/asset-status/records").status_code,401)
        self.assertEqual(self.client.post("/asset-status/records",json=SAMPLE,headers=self.headers).status_code,401)
    def test_wrong_password_and_rate_limit(self):
        for _ in range(5):
            self.assertEqual(self.client.post("/asset-status/unlock",json={"password":"wrong"},headers=self.headers).status_code,401)
        self.assertEqual(self.unlock().status_code,429)
    def test_missing_configuration_fails_closed(self):
        with patch.dict(os.environ,{"ASSET_STATUS_KEY":""}): self.assertEqual(self.unlock().status_code,503)
        with patch.dict(os.environ,{"ADMIN_PASSWORD":"admin1234"}): self.assertEqual(self.unlock().status_code,200)
    def test_origin_and_header(self):
        self.assertEqual(self.client.post("/asset-status/unlock",json={}).status_code,403)
        self.assertEqual(self.client.post("/asset-status/unlock",json={},headers={**self.headers,"Origin":"https://untrusted.example"}).status_code,403)
    def test_encrypted_roundtrip_and_revision(self):
        response = self.unlock()
        self.assertIn("HttpOnly",response.headers["Set-Cookie"])
        self.assertIn("SameSite=Strict",response.headers["Set-Cookie"])
        first = self.client.post("/asset-status/records",json=SAMPLE,headers=self.headers)
        self.assertEqual(first.status_code,200)
        stored = self.db.execute("SELECT content FROM tab_notes WHERE tab=?",(PRIVATE_TAB,)).fetchone()[0]
        self.assertNotIn("001-000",stored)
        self.assertNotIn("가상",stored)
        self.assertEqual(self.client.get("/asset-status/records").status_code,401)
        read = self.client.get("/asset-status/records",headers=self.headers).json["document"]
        self.assertEqual(validate_document(read),validate_document(SAMPLE))
        self.assertEqual(self.client.post("/asset-status/records",json=SAMPLE,headers=self.headers).status_code,409)
        read["accounts"][0]["amount"] = 200000
        self.assertEqual(self.client.put("/asset-status/records",json=read,headers=self.headers).status_code,200)
        self.assertEqual(self.client.put("/asset-status/records",json=read,headers=self.headers).status_code,409)
    def test_lock_and_revisit(self):
        self.unlock(); self.client.post("/asset-status/lock",json={},headers=self.headers)
        self.assertEqual(self.client.get("/asset-status/records").status_code,401)
        self.unlock(); response = self.client.get("/asset-status/")
        self.assertEqual(response.status_code,200)
        self.assertIn("no-store", response.headers["Cache-Control"])
        self.assertIn("default-src 'none'",response.headers["Content-Security-Policy"])
        self.assertEqual(self.client.get("/asset-status/records").status_code,401)
    def test_invalid_values(self):
        for field,value in [("amount",float("nan")),("amount","1,000"),("maturity","2026-02-30"),("account",1234)]:
            bad = copy.deepcopy(SAMPLE); bad["accounts"][0][field] = value
            with self.assertRaises(ValueError): validate_document(bad)
    def test_key_loss_does_not_overwrite_records(self):
        self.unlock(); self.client.post("/asset-status/records",json=SAMPLE,headers=self.headers)
        with patch.dict(os.environ,{"ASSET_STATUS_KEY":Fernet.generate_key().decode()}):
            self.unlock()
            self.assertEqual(self.client.post("/asset-status/records",json=SAMPLE,headers=self.headers).status_code,401)

    def test_public_notes_cannot_read_or_modify_private_storage(self):
        # Load route definitions without running existing production startup migrations.
        tree = ast.parse(Path("app.py").read_text())
        tree.body = [node for node in tree.body if not (isinstance(node,ast.Expr)
            and isinstance(node.value,ast.Call) and isinstance(node.value.func,ast.Name)
            and node.value.func.id in {"init_db","clean_snapshot_outliers"})]
        namespace = {"__name__":"asset_routes_test"}
        exec(compile(tree,"app.py","exec"),namespace)
        class Connection:
            def commit(inner): self.db.commit()
            def close(inner): pass
        namespace["get_db"] = lambda: Connection()
        namespace["query"] = lambda conn,sql,args=(): self.db.execute(sql.replace("%s","?"),args)
        self.db.execute("INSERT INTO tab_notes (id,tab,content) VALUES (100,?,?)",(PRIVATE_TAB,"ciphertext"))
        client = namespace["app"].test_client()
        self.assertEqual(client.get("/api/notes",query_string={"tab":PRIVATE_TAB}).status_code,404)
        self.assertEqual(client.post("/api/notes",json={"tab":PRIVATE_TAB,"content":"attack"}).status_code,404)
        for nid, tab in [(100, PRIVATE_TAB), (1, AUTH_TAB)]:
            original = self.db.execute("SELECT content FROM tab_notes WHERE id=?",(nid,)).fetchone()[0]
            self.assertEqual(client.get("/api/notes",query_string={"tab":tab}).status_code,404)
            self.assertEqual(client.post(f"/api/notes/{nid}/pin",json={}).status_code,404)
            client.put(f"/api/notes/{nid}",json={"content":"attack"})
            client.delete(f"/api/notes/{nid}")
            self.assertEqual(self.db.execute("SELECT content FROM tab_notes WHERE id=?",(nid,)).fetchone()[0],original)

    def test_owner_only_setup_and_no_reset(self):
        headers = {**self.headers,"X-Asset-Setup":setup_capability()}
        self.assertEqual(self.client.post("/asset-status/setup",json={"password":"replacement-password"},headers=self.headers).status_code,403)
        self.assertEqual(self.client.post("/asset-status/setup",json={"password":"replacement-password"},headers=headers).status_code,409)
        self.assertEqual(self.unlock().status_code,200)
        stored = self.db.execute("SELECT content FROM tab_notes WHERE tab=?",(AUTH_TAB,)).fetchone()[0]
        self.assertNotIn("synthetic-test-only-password",stored)

    def test_owner_setup_access_without_url_secrets(self):
        self.db.execute("DELETE FROM tab_notes WHERE tab=?",(AUTH_TAB,));self.db.commit()
        self.assertEqual(self.client.post("/asset-status/setup-access",json={},headers=self.headers).status_code,403)
        headers={**self.headers,"X-Asset-Setup":setup_capability()}
        self.assertEqual(self.client.post("/asset-status/setup-access",json={},headers=headers).status_code,200)
        self.assertEqual(self.client.post("/asset-status/setup",json={"password":"short"},headers=headers).status_code,400)
        self.assertEqual(self.client.post("/asset-status/setup",json={"password":"synthetic-test-only-password"},headers=headers).status_code,200)
        self.assertEqual(self.client.post("/asset-status/setup-access",json={},headers=headers).status_code,409)

    def test_initial_import_is_owner_only_and_cannot_overwrite_or_read(self):
        self.assertEqual(self.client.post("/asset-status/initial-import",json=SAMPLE,headers=self.headers).status_code,403)
        headers = {**self.headers,"X-Asset-Setup":setup_capability()}
        imported = self.client.post("/asset-status/initial-import",json=SAMPLE,headers=headers)
        self.assertEqual(imported.status_code,200)
        self.assertNotIn("document", imported.json)
        self.assertEqual(len(imported.json["digest"]),64)
        self.assertEqual(self.client.post("/asset-status/initial-import",json=SAMPLE,headers=headers).status_code,409)
        self.assertEqual(self.client.get("/asset-status/records",headers=headers).status_code,401)
        self.unlock()
        self.assertEqual(validate_document(self.client.get("/asset-status/records",headers=self.headers).json["document"]),validate_document(SAMPLE))

    def test_browser_workbook_import_requires_password_and_is_initial_only(self):
        self.assertEqual(self.client.post("/asset-status/initial-workbook",data=b"sample",headers=self.headers).status_code,401)
        self.unlock()
        with patch("import_asset_status.read_document",return_value=SAMPLE) as reader:
            response=self.client.post("/asset-status/initial-workbook",data=b"sample",headers=self.headers,content_type="application/octet-stream")
            self.assertEqual(response.status_code,200)
            self.assertEqual(reader.call_args.args[0].getvalue(),b"sample")
            reader.reset_mock()
            self.assertEqual(self.client.post("/asset-status/initial-workbook",data=b"replacement",headers=self.headers).status_code,409)
            reader.assert_not_called()

    def test_export_requires_password_and_preserves_safe_cell_types(self):
        self.assertEqual(self.client.post("/asset-status/export",headers=self.headers).status_code,401)
        self.unlock()
        sample=copy.deepcopy(SAMPLE)
        sample["accounts"][0]["description"]="=1+1"
        self.client.post("/asset-status/records",json=sample,headers=self.headers)
        response=self.client.post("/asset-status/export",headers=self.headers)
        self.assertEqual(response.status_code,200)
        self.assertIn("attachment",response.headers["Content-Disposition"])
        self.assertIn("no-store",response.headers["Cache-Control"])
        book=openpyxl.load_workbook(BytesIO(response.data))
        self.assertEqual(book.sheetnames,["전체자산","이전대상채권","자금이동대상","미래충족금액","소모비용"])
        self.assertEqual(book["전체자산"]["C4"].data_type,"s")
        self.assertEqual(book["전체자산"]["D4"].value,"001-000-000")
        self.assertEqual(book["전체자산"]["E4"].value,100000)
        book.close()

    def test_export_bond_completion_date_boundary(self):
        from asset_export import export_workbook
        data=validate_document(SAMPLE)
        template=copy.deepcopy(data["accounts"][0]);data["accounts"]=[]
        for i, maturity in enumerate(("2026-09-07","2026-09-08","2026-09-09","2025-01-01","")):
            row=copy.deepcopy(template);row.update(id=f"owned-{i}",group="total",year=None,maturity=maturity)
            data["accounts"].append(row)
        legacy=copy.deepcopy(template);legacy.update(id="legacy-copy",group="bonds",year=2026)
        data["accounts"].append(legacy)
        book=openpyxl.load_workbook(export_workbook(data,date(2026,9,8)))
        statuses=[book["이전대상채권"].cell(r,1).value for r in range(4,9)]
        self.assertEqual(statuses.count("이전완료"),2)
        self.assertEqual(statuses.count("이전대기"),1)
        self.assertEqual(statuses.count("날짜 확인 필요"),2)
        self.assertEqual(book["이전대상채권"]["B1"].value,5*(100000+1234))
        book.close()

    def test_linked_movement_uses_source_without_changing_total_or_review(self):
        sample=copy.deepcopy(SAMPLE)
        linked=copy.deepcopy(sample["accounts"][0]);linked.update(id="move-linked",group="movement",source_id="sample-1",amount=1,review_notes="가상 검토 메모")
        sample["accounts"].append(linked)
        clean=validate_document(sample)
        self.assertEqual(clean["accounts"][1]["amount"],100000)
        self.assertEqual(clean["accounts"][1]["review_notes"],"가상 검토 메모")
        sample["accounts"][0]["amount"]=250000
        clean=validate_document(sample)
        self.assertEqual(clean["accounts"][1]["amount"],250000)
        self.assertEqual(sum(r["amount"] for r in clean["accounts"] if r["group"]=="total"),250000)
        broken=copy.deepcopy(sample);broken["accounts"][1]["source_id"]="missing"
        with self.assertRaises(ValueError):validate_document(broken)
        duplicate=copy.deepcopy(linked);duplicate["id"]="duplicate";sample["accounts"].append(duplicate)
        with self.assertRaises(ValueError):validate_document(sample)

    def test_consumed_assets_keep_original_total_and_have_separate_balance(self):
        from asset_export import export_workbook
        data=validate_document(SAMPLE)
        data["accounts"][0]["maturity"]="2026-05-31"
        for row_id, category, maturity in (("cutoff","예적금","2026-06-01"),("coin","코인","2026-01-01")):
            row=copy.deepcopy(data["accounts"][0]);row.update(id=row_id,category=category,maturity=maturity)
            data["accounts"].append(row)
        book=openpyxl.load_workbook(export_workbook(data,date(2026,9,8)))
        ws=book["전체자산"]
        self.assertEqual(ws["B1"].value,300000)
        self.assertEqual([ws.cell(r,10).value for r in range(4,7)],["이미 소비한 자산","보유자산","보유자산"])
        self.assertEqual(ws.cell(ws.max_row,1).value,"보유자산 (자산현황)")
        self.assertEqual(ws.cell(ws.max_row,2).value,200000)
        book.close()

    def test_independent_amount_and_destination_roundtrip(self):
        sample=copy.deepcopy(SAMPLE)
        second=copy.deepcopy(sample["accounts"][0]);second.update(id="independent",amount=222000)
        sample["accounts"].append(second)
        sample["accounts"][0].update(amount=111000,destination="가상 이동처",maturity="2026-05-31")
        self.unlock()
        saved=self.client.post("/asset-status/records",json=sample,headers=self.headers)
        self.assertEqual(saved.status_code,200)
        rows=self.client.get("/asset-status/records",headers=self.headers).json["document"]["accounts"]
        self.assertEqual([r["amount"] for r in rows],[111000,222000])
        self.assertEqual(rows[0]["destination"],"가상 이동처")
        output=self.client.post("/asset-status/export",headers=self.headers)
        book=openpyxl.load_workbook(BytesIO(output.data))
        self.assertEqual(book["전체자산"]["I3"].value,"이동처")
        self.assertEqual(book["전체자산"]["I4"].value,"가상 이동처")
        book.close()


if __name__ == "__main__": unittest.main()
