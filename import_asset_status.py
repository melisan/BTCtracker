"""One-time in-memory migration. Never writes workbook data to disk or stdout."""
import argparse
import getpass
import math
import re
import uuid
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import requests


def text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        return f"{Decimal(str(value)) * 100:f}%" if "%" in cell.number_format else f"{value:,}"
    return str(value)


def read_document(path, total_scope="c1", name_colors=None):
    # Read in memory only; do not execute formulas or follow workbook links.
    book = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    formulas = openpyxl.load_workbook(path, data_only=False, keep_links=False)
    result = {"accounts":[], "other_assets":[], "future":[], "expenses":[], "other":"", "total_scope":"c1"}
    result["name_colors"] = name_colors or {}
    try:
        if len(book.worksheets) != 1:
            raise ValueError("Source structure changed")
        sheet, source = book.active, formulas.active
        if source["C1"].value.upper().replace(" ","") != "=SUM(F3:F10,F12:F13)":
            raise ValueError("Total definition changed")
        def numeric(address):
            value = sheet[address].value
            if value is None and source[address].data_type == "f":
                raise ValueError("Missing saved formula result")
            if value is not None and (isinstance(value,bool) or not isinstance(value,(int,float))):
                raise ValueError("Invalid monetary cell")
            return value
        def words(*addresses):
            return "\n".join(text(sheet[a]) for a in addresses if sheet[a].value is not None)
        for group, rows in (("total",range(3,11)),("bonds",range(18,23)),("movement",range(26,29))):
            for r in rows:
                if str(sheet[f"B{r}"].value or "").replace("·","").strip() != "예적금":
                    raise ValueError("Source row changed")
                maturity = text(sheet[f"I{r}"]).strip()
                if maturity:
                    match = re.fullmatch(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})\.?",maturity)
                    if not match: raise ValueError("Invalid maturity date")
                    maturity = date(*map(int,match.groups())).isoformat()
                result["accounts"].append({"id":uuid.uuid4().hex,"group":group,"year":2025 if group=="bonds" else None,"in_total":group=="total",
                    "category":text(sheet[f"B{r}"]),"name":text(sheet[f"C{r}"]),"description":text(sheet[f"D{r}"]),
                    "account":text(sheet[f"E{r}"]),"amount":numeric(f"F{r}"),"interest":numeric(f"J{r}"),
                    "notes":words(f"G{r}",f"H{r}"),"maturity":maturity})
                if group == "movement":
                    item = result["accounts"][-1]
                    item["notes"] = "\n".join(part for part in (
                        f"항목: {item['category']}", f"이름: {item['name']}", f"계좌번호: {item['account']}",
                        f"이자: {text(sheet[f'J{r}'])}", item["notes"], f"만기일: {maturity}" if maturity else "") if part)
        for r in (12,13):
            source_type = words(f"B{r}",f"D{r}").upper()
            if not any(v in source_type for v in ("BTC","ETH","코인","비트","이더")):
                raise ValueError("Asset category needs review")
            result["accounts"].append({"id":uuid.uuid4().hex,"group":"total","year":None,"in_total":True,"category":"코인",
                "name":words(f"C{r}"),"description":words(f"B{r}",f"D{r}"),"account":"","amount":numeric(f"F{r}"),"interest":None,
                "notes":words(f"G{r}",f"H{r}"),"maturity":""})
        expense_rows = [r for r in range(43,sheet.max_row+1) if any(sheet.cell(r,c).value is not None for c in range(2,6))]
        for collection,rows in (("future",range(34,37)),("expenses",expense_rows)):
            for r in rows:
                result[collection].append({"id":uuid.uuid4().hex,"in_total":False,
                    "description":words(f"B{r}"),"amount":numeric(f"C{r}"),"notes":words(f"D{r}",f"E{r}")})
        if not any(row["description"].strip()=="기타" for row in result["expenses"]):
            result["expenses"].append({"id":uuid.uuid4().hex,"in_total":False,"description":"기타","amount":None,"notes":""})
        result.update(bond_notes=words("C15","D15","B16"),movement_notes="",
            future_balance_label=words("B39"),future_balance_notes=words("D39"),
            future_extra_label=words("F39"),future_extra_amount=numeric("G39"),future_extra_notes=words("H39"))
        totals = {g:sum((r["amount"] or 0) for r in result["accounts"] if r["group"]==g) for g in ("total","bonds","movement")}
        interests = {g:sum((r["interest"] or 0) for r in result["accounts"] if r["group"]==g) for g in totals}
        total = totals["total"] + sum(r["amount"] or 0 for r in result["other_assets"])
        future = sum(r["amount"] or 0 for r in result["future"])
        checks = {"C1":total,"C16":totals["bonds"],"D16":totals["bonds"]+interests["bonds"],
            "F23":totals["bonds"],"J23":interests["bonds"],"K23":totals["bonds"]+interests["bonds"],
            "F29":totals["movement"],"J29":interests["movement"],"K29":totals["movement"]+interests["movement"],
            "C37":future,"C39":totals["bonds"]+interests["bonds"]-future}
        if any(numeric(cell) is None or not math.isclose(value,numeric(cell),rel_tol=1e-12,abs_tol=0.01) for cell,value in checks.items()):
            raise ValueError("Total reconciliation failed")
        known_rows = {1,12,13,15,16,23,25,29,33,37,39,42}|set(range(3,11))|set(range(18,23))|set(range(26,29))|set(range(34,37))|set(expense_rows)
        if any(c.value is not None and (c.row not in known_rows or c.column>11) for row in sheet for c in row):
            raise ValueError("Unmapped content")
        return result
    finally:
        book.close()
        formulas.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--url")
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--configure-colors", action="store_true")
    parser.add_argument("--total-scope", choices=["c1","all"], default="c1")
    args = parser.parse_args()
    try:
        from asset_status import validate_document
        data = validate_document(read_document(args.workbook,args.total_scope))
        if args.check_only:
            print("Structure and field validation passed; no data saved or transmitted.")
        else:
            if args.configure_colors:
                colors = {}
                for color in ("blue","green"):
                    name = getpass.getpass(f"Name for {color} (hidden): ").strip()
                    if name: colors[name] = color
                data["name_colors"] = colors
            target = urlparse(args.url or "")
            if target.scheme != "https" or target.netloc != "btctracker-production.up.railway.app":
                raise ValueError("Unexpected destination")
            base = f"https://{target.netloc}/asset-status"
            with requests.Session() as session:
                session.headers.update({"X-Asset-Request":"1"})
                password = getpass.getpass("App password (hidden): ")
                response = session.post(base + "/unlock", json={"password":password}, timeout=30, allow_redirects=False)
                password = None
                if response.status_code != 200:
                    raise ValueError("Authentication failed")
                session.headers["X-Asset-Token"] = response.json()["token"]
                try:
                    response = session.post(base + "/records", json=data, timeout=30, allow_redirects=False)
                    if response.status_code != 200:
                        raise ValueError("Import failed")
                    saved = session.get(base + "/records", timeout=30, allow_redirects=False)
                    if saved.status_code != 200 or validate_document(saved.json()["document"]) != data:
                        raise ValueError("Verification failed")
                finally:
                    session.post(base + "/lock", json={}, timeout=30, allow_redirects=False)
                print("Import and verification complete. No local data copies were created.")
    except Exception:
        # Exception messages can contain cells, paths, credentials, or response bodies.
        print("Operation not completed. Review configuration or source field formats locally.")
        raise SystemExit(1)
