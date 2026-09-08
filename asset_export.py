"""Build a download in memory; never persist decrypted records or workbooks."""
from io import BytesIO
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def export_workbook(data, today):
    year = today.year
    book = Workbook()
    book.remove(book.active)
    colors = data.get("name_colors", {})
    def sheet(title, headers, rows, total, note=""):
        ws = book.create_sheet(title)
        ws.append([title, total])
        ws.append([note])
        ws.append(headers)
        for values in rows:
            ws.append(values)
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # User text must never become an Excel formula.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);–'
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                    if cell.value.year == year:
                        cell.font = Font(color="C00000", bold=True)
                if isinstance(cell.value, str) and cell.value in colors:
                    cell.font = Font(color={"blue":"1565C0", "green":"228B22"}[colors[cell.value]], bold=True)
        for cell in ws[3]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="4A3B28")
        ws["A1"].font = Font(size=16, bold=True)
        ws["B1"].font = Font(size=16, bold=True)
        for col, label in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col)].width = 36 if "내용" in label else 24 if "계좌" in label else 19
            if label == "연도":
                for row in range(4, ws.max_row+1):
                    ws.cell(row,col).number_format = "0"
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(ws.max_row,3)}"
        ws.print_title_rows = "1:3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return ws
    accounts = data["accounts"]
    rows_for = lambda group: [r for r in accounts if r["group"] == group]
    amount = lambda rows, field="amount": sum(r.get(field) or 0 for r in rows)
    total, bonds, movement = (rows_for(g) for g in ("total", "bonds", "movement"))
    fields = ["category", "name", "description", "account", "amount", "interest", "notes", "maturity"]
    headers = ["항목", "이름", "내용", "계좌번호", "원본총액", "이자", "내용(비고)", "만기일"]
    def account_values(row):
        return [date.fromisoformat(row[f]) if f == "maturity" and row.get(f) else row.get(f) for f in fields]
    ws = sheet("전체자산", headers, [account_values(r) for r in total], amount(total)+amount(data.get("other_assets", [])))
    for r in data.get("other_assets", []):
        ws.append(["기타", "", r["description"], "", r["amount"], None, r["notes"]])
        for c in ws[ws.max_row]:
            if isinstance(c.value,str): c.data_type="s"
    completed = lambda r: bool(r.get("maturity", "").startswith("2026-") and r["maturity"] < today.isoformat())
    sheet("이전대상채권", ["상태","연도"]+headers, [["이전완료" if completed(r) else "이전대기",r["year"]]+account_values(r) for r in sorted(bonds,key=lambda r:(not completed(r),r["year"]))], amount(bonds)+amount(bonds,"interest"), data.get("bond_notes", ""))
    sheet("자금이동대상", ["항목","금액","내용"], [[r["description"],r["amount"],r["notes"]] for r in movement], amount(movement), data.get("movement_notes", ""))
    future = data.get("future", [])
    available = amount(bonds)+amount(bonds,"interest")
    future_rows = [[r["description"],r["amount"],r["notes"]] for r in future]
    future_rows += [["정기예금금액총액",available,""],["채무청산금",amount(future),""],["유용금액",available-amount(future),"※ 유용금액= 정기예금금액총액- 채무청산금"]]
    if data.get("future_extra_label") or data.get("future_extra_amount") is not None:
        future_rows.append([data.get("future_extra_label",""),data.get("future_extra_amount"),data.get("future_extra_notes","")])
    sheet("미래충족금액", ["항목","금액","내용"], future_rows, amount(future), data.get("future_balance_notes", ""))
    expenses = data.get("expenses", [])
    sheet("소모비용", ["항목","금액총액","내용"], [[r["description"],r["amount"],r["notes"]] for r in expenses], amount(expenses), data.get("other", ""))
    output = BytesIO()
    book.save(output)
    book.close()
    output.seek(0)
    return output
